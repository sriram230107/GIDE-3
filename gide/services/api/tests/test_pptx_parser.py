import uuid
from pptx import Presentation
from pptx.util import Inches
from app.ingestion.pptx_parser import extract_slide_text, extract_pptx


def _deck(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Gradient Descent"
    s.placeholders[1].text = "Move against the gradient\nlearning rate controls step size"
    s.notes_slide.notes_text_frame.text = "Mention overshooting."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Backprop"
    t = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    t.cell(0, 0).text, t.cell(0, 1).text = "layer", "grad"
    p = tmp_path / "deck.pptx"
    prs.save(p)
    return str(p)


def test_text_notes_tables(tmp_path):
    slides = extract_slide_text(_deck(tmp_path))
    assert [s["slide_number"] for s in slides] == [1, 2]
    assert "learning rate" in slides[0]["body"] and slides[0]["notes"] == "Mention overshooting."
    assert "layer | grad" in slides[1]["body"]


def test_units_have_slide_provenance_and_report_missing_render(tmp_path):
    units, meta = extract_pptx(_deck(tmp_path), uuid.uuid4(), None)
    assert all(u["slide_number"] in (1, 2) for u in units)
    assert any(u["unit_metadata"].get("is_speaker_note") for u in units)
    assert meta["slide_count"] == 2 and meta["slides_rendered"] == 0
